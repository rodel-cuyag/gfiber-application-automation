"""
prior_day.py
-------------
Best-effort lookup of the previous day's already-generated EOD Report
workbook, so the current run's "Yesterday" column can be filled in
without recomputing anything from the source data.

One dashboard row ("System Errors") is written as a live Excel formula
rather than a literal value (see excel_writer.py) — openpyxl can't
evaluate a formula's result, only Excel/LibreOffice can, on open. That
row is simply skipped here and left for the caller to treat as "no data
available", same as a missing prior report.
"""

from openpyxl import load_workbook

from src import config
from src.excel_writer import BANNER_SEP


def _find_previous_report(agent_id, previous_date):
    """
    Returns the most recently modified EOD report file for *previous_date*
    (searching recursively across that date's run-time subfolders, and
    covering any `resolve_output_path` rerun suffixes like " (1).xlsx"
    within a given run), or None if no prior report exists.
    """
    eod_dir = config.get_eod_output_dir(previous_date, previous_date)
    if not eod_dir.exists():
        return None

    stem = config.OUTPUT_FILENAME_TEMPLATE_SINGLE.format(
        agent_id=agent_id, start_date=previous_date,
    ).rsplit(".xlsx", 1)[0]
    candidates = list(eod_dir.rglob(f"{stem}*.xlsx"))
    if not candidates:
        return None
    return max(candidates, key=lambda p: p.stat().st_mtime)


def _band_titles_by_row(ws) -> dict:
    """
    Returns {row_number: band title} for every full-width A:D band — the
    navy section dividers and the breakdown banners written by
    excel_writer._band.

    Detection is by merge geometry rather than cell contents: a manually
    filled row like "Open P0 issues" also has a non-empty Column A and an
    empty Column B, so sniffing the values would misclassify it as a band.
    """
    titles = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_col == 1 and rng.max_col == 4 and rng.min_row == rng.max_row:
            text = ws.cell(row=rng.min_row, column=1).value
            if text is None:
                continue
            # Banners carry their parent's total ("... Breakdown — 57"),
            # which changes daily; key off the stable title only.
            titles[rng.min_row] = str(text).split(BANNER_SEP)[0].strip()
    return titles


def load_previous_day_values(agent_id, previous_date) -> dict:
    """
    Returns {(band_title, metric_label): value} read from Column A/B of the
    previous day's "EOD Report" sheet, for every row with a literal
    (non-formula) value in Column B. Returns {} if no prior report is found
    or it can't be read.

    The key is composite because display labels repeat across the sheet —
    "Wishes to Proceed" appears under the Postpaid and Non-Postpaid
    breakdown banners as well as in OUTCOMES. Keying on the label alone
    would collapse all three onto whichever was read last. Rows above the
    first band use "" as their band title, matching the writer.
    """
    path = _find_previous_report(agent_id, previous_date)
    if path is None:
        return {}

    try:
        wb = load_workbook(path, data_only=False)
        if "EOD Report" not in wb.sheetnames:
            return {}
        ws = wb["EOD Report"]

        band_titles = _band_titles_by_row(ws)
        group = ""
        values = {}
        for row in ws.iter_rows(min_row=5, max_col=2):
            label_cell, value_cell = row[0], row[1]
            if label_cell.row in band_titles:
                group = band_titles[label_cell.row]
                continue
            if label_cell.value is None:
                continue  # blank row
            value = value_cell.value
            if isinstance(value, str) and value.startswith("="):
                continue  # live formula; can't reliably read its result
            values[(group, label_cell.value)] = value
        return values
    except Exception:
        return {}
